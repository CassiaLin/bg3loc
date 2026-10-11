"""Public, complete synthetic Basic state/cache regressions; never a real game."""
from pathlib import Path
import json
from unittest.mock import patch
import pytest

from bg3loc.e2e.validate import run_validate
from bg3loc.e2e.validate import WorkflowValidateError
from bg3loc.e2e.rebuild import WorkflowRebuildError
from bg3loc.e2e.install import run_workflow_install, WorkflowInstallError
from bg3loc.e2e.state import state_path
from test_basic_baseline_integrity import make_basic, rebuild, sha, write_rows
from test_locale_material_integrity import backend_scope, json_rows
from test_legacy_pak_integrity import ZipBackend
from test_source_package_freshness import replace_valid_package
from openpyxl import load_workbook

def edit_any_return(b,text):
    if b['returns'].suffix=='.jsonl':
        edit_return(b,text)
    else:
        book=load_workbook(b['returns']);sheet=book.active;fields={c.value:i+1 for i,c in enumerate(sheet[1])}
        for index in range(2,sheet.max_row+1):
            if sheet.cell(index,fields['ContentUid']).value=='fictional-shared':sheet.cell(index,fields['ProposedTargetText']).value=text
        book.save(b['returns']);book.close()

@pytest.fixture(autouse=True)
def archive_backend():
    with backend_scope():yield

@pytest.fixture
def basic(tmp_path):return make_basic(tmp_path)

def edit_return(b,text):
    rows=json_rows(b['returns'])
    for row in rows:
        if row['ContentUid']=='fictional-shared':row['ProposedTargetText']=text
    write_rows(b['returns'],rows)

def current_state(b):return json.loads(state_path(b['config']).read_text())

def test_new_validation_fail_revokes_current_install_eligibility(basic):
    b=basic;result=rebuild(b);old=current_state(b)['artifacts']['rebuildManifest']
    edit_return(b,'SYNTHETIC missing runtime token')
    assert run_validate(b['config'])['status']=='fail'
    with pytest.raises(WorkflowInstallError):run_workflow_install(b['config'],apply=False)
    assert Path(old).is_file(), 'A historical successful manifest must be retained'

def test_cached_dry_run_rechecks_changed_artifact(basic):
    b=basic;result=rebuild(b);run_workflow_install(b['config'],apply=False)
    artifact=Path(next(a['path'] for a in result['artifacts'] if a['type']=='pak'))
    artifact.write_bytes(b'SYNTHETIC tampered candidate')
    with pytest.raises(WorkflowInstallError):run_workflow_install(b['config'],apply=False)

def test_cached_dry_run_rechecks_changed_target(basic):
    b=basic;rebuild(b);run_workflow_install(b['config'],apply=False)
    replace_valid_package(b['game']/'Data/Localization/French.pak')
    with pytest.raises(WorkflowInstallError):run_workflow_install(b['config'],apply=False)

def test_current_install_rejects_returns_edited_without_validation(basic):
    b=basic;rebuild(b);edit_return(b,'SYNTHETIC later edit {name}.')
    with pytest.raises(WorkflowInstallError):run_workflow_install(b['config'],apply=False)

def test_current_rebuild_rejects_returns_edited_without_validation(basic):
    b=basic;edit_return(b,'SYNTHETIC later edit {name}.')
    with pytest.raises(WorkflowRebuildError):rebuild(b)

def test_failed_rebuild_keeps_all_prior_success_bytes_isolated(basic):
    b=basic;first=rebuild(b)
    old_manifest=Path(current_state(b)['artifacts']['rebuildManifest'])
    before={p:p.read_bytes() for p in [old_manifest,*[Path(a['path']) for a in first['artifacts']]]}
    edit_return(b,'SYNTHETIC valid revision {name}.');assert run_validate(b['config'])['status']=='pass'
    class Broken(ZipBackend):
        def create_package(self,*args):raise OSError('SYNTHETIC repack failure')
    with pytest.raises(WorkflowRebuildError):rebuild(b,Broken())
    with pytest.raises(WorkflowInstallError):run_workflow_install(b['config'],apply=False)
    assert all(p.read_bytes()==raw for p,raw in before.items()), 'Failure must not overwrite prior successful LOCA/manifest/PAK'

def test_same_input_legal_validation_reuses_current_rebuild_eligibility(basic):
    b=basic;rebuild(b)
    assert run_validate(b['config']).get('reused') is True
    assert run_workflow_install(b['config'],apply=False)['status']=='dry-run-pass'

def test_cached_dry_run_checks_current_validation_state(basic):
    b=basic;rebuild(b);run_workflow_install(b['config'],apply=False)
    state=current_state(b);state['stages']['validate']['status']='running';state_path(b['config']).write_text(json.dumps(state))
    with pytest.raises(WorkflowInstallError):run_workflow_install(b['config'],apply=False)

@pytest.mark.parametrize('fmt',['jsonl','xlsx'])
def test_validation_and_rebuild_pass_fail_pass_recovery(tmp_path,fmt):
    b=make_basic(tmp_path,fmt);first=rebuild(b);first_state=current_state(b)
    first_manifest=Path(first_state['artifacts']['rebuildManifest']);first_bytes=first_manifest.read_bytes()
    run_workflow_install(b['config'],apply=False)
    edit_any_return(b,'SYNTHETIC missing token')
    assert run_validate(b['config'])['status']=='fail'
    assert current_state(b)['stages']['rebuild']['status']!='completed'
    with pytest.raises(WorkflowInstallError):run_workflow_install(b['config'],apply=False)
    edit_any_return(b,'SYNTHETIC good revision {name}.');assert run_validate(b['config'])['status']=='pass'
    class Broken(ZipBackend):
        def create_package(self,*args):raise OSError('SYNTHETIC failure')
    with pytest.raises(WorkflowRebuildError):rebuild(b,Broken())
    assert current_state(b)['stages']['rebuild']['status']=='failed'
    assert first_manifest.read_bytes()==first_bytes
    assert current_state(b)['artifacts']['rebuildManifest']==str(first_manifest)
    failed=b['root']/'workspace/e2e/rebuild/failed'
    assert list(failed.rglob('*.loca')) and not list(failed.rglob('rebuild-manifest.json'))
    second=rebuild(b);second_manifest=Path(current_state(b)['artifacts']['rebuildManifest'])
    assert second_manifest!=first_manifest and first_manifest.read_bytes()==first_bytes
    assert run_workflow_install(b['config'],apply=False)['status']=='dry-run-pass'

def test_unchanged_rebuild_reuses_verified_artifacts_but_dry_run_preflight_is_fresh(basic):
    b=basic;first=rebuild(b);old=current_state(b)
    with patch('bg3loc.e2e.rebuild.run_rebuild',side_effect=AssertionError('unnecessary rebuild')):
        assert rebuild(b)==first
    import bg3loc.e2e.install as module
    actual=module.run_install
    with patch.object(module,'run_install',wraps=actual) as preflight:
        a=run_workflow_install(b['config'],apply=False);z=run_workflow_install(b['config'],apply=False)
    assert a==z and preflight.call_count==2
    assert current_state(b)['artifacts']['rebuildManifest']==old['artifacts']['rebuildManifest']

@pytest.mark.parametrize('attack',['accepted','validate-manifest','validate-fingerprint','rebuild-state','rebuild-manifest','return-added','return-missing'])
def test_all_current_bindings_are_checked_before_install(basic,attack):
    b=basic;rebuild(b);run_workflow_install(b['config'],apply=False);state=current_state(b)
    if attack=='accepted':
        validate=json.loads(Path(state['artifacts']['validateManifest']).read_text());Path(validate['accepted']['path']).write_text('')
    elif attack=='validate-manifest':Path(state['artifacts']['validateManifest']).write_text('{}')
    elif attack=='validate-fingerprint':
        state['stages']['validate']['fingerprint']='0'*64;state_path(b['config']).write_text(json.dumps(state))
    elif attack=='rebuild-state':
        state['stages']['rebuild']['fingerprint']='0'*64;state_path(b['config']).write_text(json.dumps(state))
    elif attack=='rebuild-manifest':Path(state['artifacts']['rebuildManifest']).write_text('{}')
    elif attack=='return-added':(b['returns'].parent/'extra.jsonl').write_bytes(b['returns'].read_bytes())
    else:b['returns'].unlink()
    with patch('bg3loc.e2e.install.run_install',side_effect=AssertionError('must not transport a stale install')):
        with pytest.raises(WorkflowInstallError):run_workflow_install(b['config'],apply=False)

def test_explicit_external_input_location_is_preserved_and_checked(basic):
    b=basic;external=b['root']/'external-return.jsonl';external.write_bytes(b['returns'].read_bytes())
    assert run_validate(b['config'],str(external))['status']=='pass'
    assert current_state(b)['stages']['translate']['artifact']==str(external.resolve())
    rebuild(b)
    # Unused default returns cannot silently become the translation provenance.
    edit_return(b,'SYNTHETIC default-only edit')
    assert run_workflow_install(b['config'],apply=False)['status']=='dry-run-pass'
    external.write_text(external.read_text()+'\n')
    with pytest.raises(WorkflowInstallError):run_workflow_install(b['config'],apply=False)

def test_validation_start_and_parser_exception_revoke_old_eligibility(basic):
    b=basic;rebuild(b);edit_return(b,'SYNTHETIC edited {name}.')
    def failure(*args):
        state=current_state(b)
        assert state['stages']['validate']['status']=='running'
        assert state['stages']['rebuild']['status']!='completed'
        raise ValueError('SYNTHETIC parser failure')
    with patch('bg3loc.e2e.validate.read_return_material',side_effect=failure),pytest.raises(WorkflowValidateError):run_validate(b['config'])
    assert current_state(b)['stages']['validate']['status']=='failed'
    with pytest.raises(WorkflowInstallError):run_workflow_install(b['config'],apply=False)

def test_injected_state_change_during_dry_run_is_not_published_as_pass(basic):
    b=basic;rebuild(b)
    import bg3loc.e2e.install as module
    original=module.run_install
    def mutate(request):
        result=original(request);state=current_state(b);state['stages']['validate']['status']='running';state_path(b['config']).write_text(json.dumps(state));return result
    with patch.object(module,'run_install',side_effect=mutate),pytest.raises(WorkflowInstallError):run_workflow_install(b['config'],apply=False)
    assert current_state(b)['stages']['install-dry-run']['status']=='failed'

def test_history_low_level_preflight_does_not_grant_current_workflow_eligibility(basic):
    from bg3loc.installer import InstallRequest,run_install
    b=basic;rebuild(b);state=current_state(b);history=Path(state['artifacts']['rebuildManifest'])
    edit_return(b,'SYNTHETIC missing token');assert run_validate(b['config'])['status']=='fail'
    with pytest.raises(WorkflowInstallError):run_workflow_install(b['config'],apply=False)
    assert run_install(InstallRequest(history,Path(state['artifacts']['scanManifest']),dry_run=True))['status']=='dry-run-pass'

def test_last_success_validation_snapshot_is_not_overwritten_on_fail(basic):
    b=basic;rebuild(b);state=current_state(b);manifest=Path(state['artifacts']['validateManifest']);content=manifest.read_bytes()
    accepted=Path(json.loads(content)['accepted']['path']);raw=accepted.read_bytes()
    edit_return(b,'SYNTHETIC missing token');assert run_validate(b['config'])['status']=='fail'
    assert manifest.read_bytes()==content and accepted.read_bytes()==raw
