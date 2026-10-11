"""CCR-B1-01 public complete synthetic workflows. Source gate is NEVER mocked."""
from contextlib import nullcontext
from pathlib import Path
import json
import os
import hashlib
import zipfile
from unittest.mock import patch
import pytest
from openpyxl import load_workbook

from bg3loc.cli import main
from bg3loc.e2e.baseline import BasicSourceSnapshot, RECEIPT_PATH, validated_baseline
from bg3loc.e2e.config import load_project_config
from bg3loc.e2e.prepare import run_prepare, WorkflowPrepareError
from bg3loc.e2e.validate import run_validate, WorkflowValidateError
from bg3loc.e2e.rebuild import run_workflow_rebuild, WorkflowRebuildError
from bg3loc.e2e.install import run_workflow_install, WorkflowInstallError
from bg3loc.e2e.state import state_path
from bg3loc.rebuilder import parse_loca_xml
from test_locale_material_integrity import installed_game, backend_scope, json_rows
from test_source_package_freshness import replace_valid_package
from test_source_package_freshness import nonlocking_pins
from test_legacy_pak_integrity import ZipBackend

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def write_rows(p,rows):p.write_text(''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in rows),encoding='utf-8')
@pytest.fixture(autouse=True)
def synthetic_archive_backend():
    # Backend transport only is fake; all source hashes, pins, receipts and
    # origin re-extraction/normalization validation run for real.
    with backend_scope():yield
def rebuild(b,backend=None):
    with patch('bg3loc.rebuilder.resolve_rebuild_backend',return_value=backend or ZipBackend()):return run_workflow_rebuild(b['config'])
def make_basic(root,fmt='jsonl',reference=False):
    game=installed_game(root)
    for pak in (game/'Data/Localization').glob('*.pak'):
        with zipfile.ZipFile(pak,'a') as z:
            for name,data in [('opaque.bin',b'\x00\xff'),('nested/meta.json',b'{"synthetic":true}'),('empty.bin',b'')]:z.writestr(name,data)
    project=root/'project.json'
    args=['project','init','--name','Synthetic Basic integrity','--source','english','--target','french','--game-dir',str(game),
          '--format',fmt,'--content-uid','fictional-shared','--content-uid','fictional-source-only','--content-uid','fictional-blank','--output',str(project)]
    if reference:args+=['--reference','german']
    assert main(args)==0
    config=load_project_config(project)
    with backend_scope():state=run_prepare(config)
    returns=next((root/'workspace/e2e/returns').glob('*.'+fmt))
    def target(uid):return 'Salut {name}.' if uid=='fictional-shared' else 'Texte <b>synthetique</b>.' if uid=='fictional-source-only' else ''
    if fmt=='jsonl':
        rows=json_rows(returns)
        for row in rows:
            row['ProposedTargetText']=target(row['ContentUid']);row['TranslationStatus']='intentionally-empty' if row['ContentUid']=='fictional-blank' else 'translated'
        write_rows(returns,rows)
    else:
        book=load_workbook(returns);sheet=book.active;fields={c.value:i+1 for i,c in enumerate(sheet[1])}
        for index in range(2,sheet.max_row+1):
            uid=sheet.cell(index,fields['ContentUid']).value
            sheet.cell(index,fields['ProposedTargetText']).value=target(uid)
            sheet.cell(index,fields['TranslationStatus']).value='intentionally-empty' if uid=='fictional-blank' else 'translated'
        book.save(returns);book.close()
    assert run_validate(config)['status']=='pass'
    extract=Path(state['artifacts']['extractManifest'])
    return dict(config=config,state=state,root=root,game=game,returns=returns,extract=extract,infos=json.loads(extract.read_text())['locales'])

@pytest.fixture
def basic(tmp_path):return make_basic(tmp_path)

@pytest.mark.parametrize('attack',['source-pak','target-pak','target-xml','extract-reference','source-normalized','target-loca','scan','aligned','missing-source'])
def test_original_b1_attacks_reject_before_rebuild_publication(basic,attack):
    b=basic;s=next(r for r in b['infos'] if r['localeId']=='English');t=next(r for r in b['infos'] if r['localeId']=='French')
    if attack in {'source-pak','target-pak'}:replace_valid_package(Path((s if attack=='source-pak' else t)['packageFile']))
    elif attack=='target-xml':
        p=Path(t['sourceXml']);p.write_text(p.read_text().replace('Texte conserve.','SYNTHETIC changed'),encoding='utf-8')
    elif attack=='extract-reference':
        p=Path(t['sourceXml']);alternate=p.with_name('alternate.xml');alternate.write_bytes(p.read_bytes())
        data=json.loads(b['extract'].read_text());next(r for r in data['locales'] if r['localeId']=='French')['sourceXml']=str(alternate);b['extract'].write_text(json.dumps(data))
    elif attack=='source-normalized':
        p=Path(s['normalized']);rows=json_rows(p)
        for r in rows:r['version']=99
        write_rows(p,rows)
    elif attack=='target-loca':Path(t['sourceLoca']).write_bytes(b'SYNTHETIC changed LOCA')
    elif attack=='scan':Path(b['state']['artifacts']['scanManifest']).write_text('{}')
    elif attack=='aligned':
        p=Path(json.loads(b['extract'].read_text())['aligned']);p.write_text(p.read_text()+'\n')
    else:Path(s['normalized']).unlink()
    with pytest.raises(WorkflowRebuildError,match='Scan / Extract / Prepare / Validate'):rebuild(b)
    assert not (b['root']/'workspace/e2e/rebuild').exists()
    with pytest.raises(WorkflowValidateError):run_validate(b['config'])

@pytest.mark.parametrize('fmt',['jsonl','xlsx'])
@pytest.mark.parametrize('reference',[False,True])
def test_full_basic_formats_versions_blank_subset_and_entry_hashes(tmp_path,fmt,reference):
    b=make_basic(tmp_path,fmt,reference);before={p:sha(p) for p in (b['game']/'Data/Localization').glob('*.pak')}
    result=rebuild(b);records={r.content_uid:r for r in parse_loca_xml(Path(next(a['path'] for a in result['artifacts'] if a['type']=='loca')))}
    assert records['fictional-shared'].version==7 and records['fictional-source-only'].version==2
    assert records['fictional-blank'].text=='' and records['fictional-blank'].version==4
    assert records['fictional-target-only'].text=='Texte conserve.'
    pak=Path(next(a['path'] for a in result['artifacts'] if a['type']=='pak'))
    with zipfile.ZipFile(b['game']/'Data/Localization/French.pak') as original,zipfile.ZipFile(pak) as rebuilt:
        assert set(original.namelist())==set(rebuilt.namelist())
        assert all(original.read(n)==rebuilt.read(n) for n in original.namelist() if not n.endswith('.loca'))
    assert run_workflow_install(b['config'],apply=False)['status']=='dry-run-pass'
    assert before=={p:sha(p) for p in before}
    assert result['extractManifest']==str(b['extract'])
    assert result['baseline']['sourceXml']==next(r['sourceXml'] for r in b['infos'] if r['localeId']=='French')
    assert not (b['root']/'workspace/e2e/backups').exists()
    assert not list((b['root']/'workspace').rglob('*.sqlite3'))

def modify_receipt(b,edit):
    package=Path(b['state']['artifacts']['translationPackage']);data=json.loads(package.read_text());receipt=package.parent/RECEIPT_PATH
    payload=json.loads(receipt.read_text());edit(payload);receipt.write_text(json.dumps(payload))
    next(a for a in data['internal']['bindings'] if a['path']==RECEIPT_PATH)['sha256']=sha(receipt);package.write_text(json.dumps(data))
    state=json.loads(state_path(b['config']).read_text());state['bindings']['translationPackageSha256']=sha(package);state_path(b['config']).write_text(json.dumps(state))

@pytest.mark.parametrize('attack',['version','unknown','missing-file','ordering','duplicate-role','size','wrong-sha','relative-path'])
def test_strict_receipt_even_with_outer_rebound_hashes(basic,attack):
    def edit(p):
        if attack=='version':p['contractVersion']='basic-source-baseline/99'
        elif attack=='unknown':p['untrustedExtra']=True
        elif attack=='missing-file':p['files'].pop()
        elif attack=='ordering':p['files'].reverse()
        elif attack=='duplicate-role':p['files'][-1]=p['files'][0]
        elif attack=='size':p['files'][0]['sizeBytes']=True
        elif attack=='wrong-sha':p['files'][0]['sha256']='0'*64
        else:p['files'][0]['path']='../foreign'
    modify_receipt(basic,edit)
    with pytest.raises(WorkflowRebuildError):rebuild(basic)

def test_old_package_has_no_auto_backfill_or_edited_return_overwrite(basic):
    b=basic;package=Path(b['state']['artifacts']['translationPackage']);data=json.loads(package.read_text());data['internal']['bindings']=[];package.write_text(json.dumps(data))
    state=json.loads(state_path(b['config']).read_text());state['bindings']['translationPackageSha256']=sha(package);state['stages']['package']['sha256']=sha(package);state_path(b['config']).write_text(json.dumps(state))
    before=b['returns'].read_bytes();receipt=(package.parent/RECEIPT_PATH).read_bytes()
    with pytest.raises(WorkflowRebuildError,match='receipt'):rebuild(b)
    with backend_scope(),pytest.raises(WorkflowPrepareError,match='receipt'):run_prepare(b['config'])
    assert b['returns'].read_bytes()==before and (package.parent/RECEIPT_PATH).read_bytes()==receipt

def test_install_checks_source_even_when_dry_run_was_cached(basic):
    b=basic;rebuild(b);run_workflow_install(b['config'],apply=False)
    replace_valid_package(b['game']/'Data/Localization/English.pak')
    with pytest.raises(WorkflowInstallError,match='Scan / Extract'):run_workflow_install(b['config'],apply=False)

@pytest.mark.parametrize('changed',['pak','xml','normalized'])
def test_private_rebuild_inputs_and_windows_original_write_denial(basic,changed):
    b=basic;target=next(r for r in b['infos'] if r['localeId']=='French')
    original=Path(target[{'pak':'packageFile','xml':'sourceXml','normalized':'normalized'}[changed]])
    import bg3loc.rebuilder as module
    real_parse=module.parse_loca_xml;observed=[]
    def during_read(path):
        assert Path(path)!=Path(target['sourceXml'])
        observed.append(path)
        if os.name=='nt':
            with pytest.raises(PermissionError):original.write_bytes(b'SYNTHETIC forbidden mutation')
        return real_parse(path)
    with patch.object(module,'parse_loca_xml',side_effect=during_read):rebuild(b)
    assert observed

def test_closed_snapshot_rejects_access_and_check(basic):
    with validated_baseline(basic['config'],basic['state']) as snapshot:pass
    with pytest.raises(ValueError,match='closed'):snapshot.file('xml','French')
    with pytest.raises(ValueError,match='closed'):snapshot.check()

def test_prepare_will_not_seal_mutated_xml_after_extract(tmp_path):
    game=installed_game(tmp_path);project=tmp_path/'p.json'
    assert main(['project','init','--name','synthetic late capture','--source','English','--target','French','--game-dir',str(game),'--output',str(project)])==0
    import bg3loc.e2e.prepare as module
    real=module.run_extract
    def corrupt(request):
        result=real(request);p=Path(next(r['sourceXml'] for r in result['locales'] if r['localeId']=='French'))
        p.write_text(p.read_text().replace('Texte conserve.','SYNTHETIC untrusted'),encoding='utf-8');return result
    with backend_scope(),patch.object(module,'run_extract',side_effect=corrupt),pytest.raises(WorkflowPrepareError,match='derive'):run_prepare(load_project_config(project))
    assert not (tmp_path/'workspace/e2e/package/translation-package-manifest.json').exists()

def test_b2_cached_artifact_and_failed_validation_are_now_rejected(basic):
    b=basic;result=rebuild(b);run_workflow_install(b['config'],apply=False)
    artifact=Path(next(a['path'] for a in result['artifacts'] if a['type']=='pak'));good=artifact.read_bytes();artifact.write_bytes(b'SYNTHETIC tampered artifact')
    # B1 input closure remains intact; current eligibility is supplied by B2.
    with pytest.raises(WorkflowInstallError):run_workflow_install(b['config'],apply=False)
    artifact.write_bytes(good)
    rows=json_rows(b['returns'])
    for r in rows:
        if r['ContentUid']=='fictional-shared':r['ProposedTargetText']='SYNTHETIC missing token'
    write_rows(b['returns'],rows);assert run_validate(b['config'])['status']=='fail'
    with pytest.raises(WorkflowInstallError):run_workflow_install(b['config'],apply=False)

def test_synthetic_install_pin_handoff_and_existing_rollback(basic):
    # Only the test-owned invented ZIP install root is modified. Never real BG3.
    from bg3loc.installer import run_rollback
    b=basic;destination=b['game']/'Data/Localization/French.pak';before=destination.read_bytes()
    rebuild(b);run_workflow_install(b['config'],apply=False)
    result=run_workflow_install(b['config'],apply=True)
    assert destination.read_bytes()!=before
    assert sha(destination)==result['deployment'][0]['installedSha256']
    assert run_rollback(Path(result['manifestPath']))['status']=='rolled-back'
    assert destination.read_bytes()==before

@pytest.mark.parametrize('kind',['pak','private-xml','private-manifest'])
def test_mutation_is_detected_without_windows_handles(basic,monkeypatch,kind):
    # Exercise rehash fallback; this is not native Linux/macOS acceptance.
    nonlocking_pins(monkeypatch)
    with pytest.raises(ValueError):
        with validated_baseline(basic['config'],basic['state']) as snapshot:
            if kind=='pak':replace_valid_package(basic['game']/'Data/Localization/English.pak',in_place=True)
            elif kind=='private-xml':snapshot.file('xml','French').write_bytes(b'SYNTHETIC changed private baseline')
            else:snapshot.private_extract.write_text('{}')
            snapshot.check()

def test_b1_failure_preserves_prior_rebuild_artifacts_and_edited_returns(basic):
    b=basic;rebuild(b);out=b['root']/'workspace/e2e/rebuild'
    before={p.relative_to(out):p.read_bytes() for p in out.rglob('*') if p.is_file()}
    edited=b['returns'].read_bytes()
    replace_valid_package(b['game']/'Data/Localization/English.pak')
    with pytest.raises(WorkflowRebuildError):rebuild(b)
    assert before=={p.relative_to(out):p.read_bytes() for p in out.rglob('*') if p.is_file()}
    assert b['returns'].read_bytes()==edited

def test_cannot_reseal_changed_xml_even_with_rebound_receipt_package_state(basic):
    b=basic;xml=Path(next(r['sourceXml'] for r in b['infos'] if r['localeId']=='French'))
    xml.write_text(xml.read_text().replace('Texte conserve.','SYNTHETIC forged baseline'),encoding='utf-8')
    def reseal(payload):
        row=next(f for f in payload['files'] if f['role']=='xml' and f['localeId']=='French')
        row['sizeBytes']=xml.stat().st_size;row['sha256']=sha(xml)
    modify_receipt(b,reseal)
    with backend_scope(),pytest.raises(WorkflowRebuildError,match='derive'):rebuild(b)
