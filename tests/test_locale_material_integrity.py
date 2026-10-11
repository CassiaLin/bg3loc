"""Synthetic public Scan/Extract/Build and Project/Workflow locale boundaries."""
from contextlib import contextmanager
from copy import deepcopy
from pathlib import Path
import json
import shutil
import zipfile
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from bg3loc.backends import ArchiveEntry, BackendProbe
from bg3loc.builder import BuildFailure, BuildRequest, run_build
from bg3loc.cli import main
from bg3loc.e2e.package import prepare_translation_package
from bg3loc.e2e.prepare import WorkflowPrepareError
from bg3loc.extractor import ExtractFailure, ExtractRequest, run_extract
from bg3loc.rebuilder import LocaRecord, write_loca_xml


SOURCE = [LocaRecord("fictional-shared", 1, "Hello {name}."),
          LocaRecord("fictional-source-only", 2, "Source <b>text</b>."),
          LocaRecord("fictional-blank", 3, "")]
TARGET = [LocaRecord("fictional-shared", 7, "Bonjour {name}."),
          LocaRecord("fictional-target-only", 9, "Texte conserve."),
          LocaRecord("fictional-blank", 4, "")]
REFERENCE = [LocaRecord("fictional-shared", 5, "Hallo {name}."),
             LocaRecord("fictional-reference-only", 6, "Nur Referenz.")]
CASINGS = [("English", "French", "German"), ("english", "french", "german"),
           ("eNgLiSh", "FrEnCh", "gErMaN")]


class SyntheticBackend:
    id = "synthetic-locale-audit"

    def list_archive(self, package, expression="*"):
        with zipfile.ZipFile(package) as archive:
            return [ArchiveEntry(i.filename, i.file_size, i.CRC) for i in archive.infolist()]

    def extract_single_file(self, package, entry, destination):
        destination.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(package) as archive:
            destination.write_bytes(archive.read(entry))

    def convert_loca(self, source, destination):
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, destination)


@contextmanager
def backend_scope():
    backend = SyntheticBackend()
    with patch("bg3loc.scanner.resolve_backend", return_value=BackendProbe(backend.id, True)), \
         patch("bg3loc.scanner.backend_from_probe", return_value=backend), \
         patch("bg3loc.extractor.backend_from_manifest", return_value=backend):
        yield backend


def installed_game(root):
    game = root / "fictional-game"
    packages = game / "Data/Localization"
    packages.mkdir(parents=True)
    for locale, rows in [("English", SOURCE), ("French", TARGET), ("German", REFERENCE)]:
        xml = root / (locale + ".xml")
        write_loca_xml(xml, rows)
        with zipfile.ZipFile(packages / (locale + ".pak"), "w") as archive:
            archive.writestr(f"Localization/{locale}/{locale}.loca", xml.read_bytes())
    return game


def extracted(root, casings=CASINGS[0]):
    game = installed_game(root)
    scan = root / "scan"
    output = root / "extract"
    with backend_scope():
        assert main(["scan", "--game-dir", str(game), "--output", str(scan)]) == 0
        assert main(["extract", "--scan", str(scan / "scan-manifest.json"),
                     "--source", casings[0], "--target", casings[1], "--reference", casings[2],
                     "--output", str(output)]) == 0
    return output / "extract-manifest.json", scan / "scan-manifest.json"


def json_rows(path):
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def assert_source_projection(rows):
    assert len(rows) == 5
    assert sum(row["TranslationRequired"] for row in rows) == len(SOURCE)
    by_uid = {row["ContentUid"]: row for row in rows}
    for source in SOURCE:
        row = by_uid[source.content_uid]
        assert row["SourceText"] == source.text and row["TranslationRequired"] is True
    assert by_uid["fictional-target-only"]["TranslationRequired"] is False
    assert by_uid["fictional-reference-only"]["TranslationRequired"] is False
    assert all(row["SourceLocale"] == "English" and row["TargetLocale"] == "French" for row in rows)


@pytest.mark.parametrize("casings", CASINGS)
def test_scan_extract_build_jsonl_preserves_canonical_roles_and_source(tmp_path, casings):
    path, _ = extracted(tmp_path, casings)
    extract = json.loads(path.read_text())
    manifest = run_build(BuildRequest(path, format="jsonl", output=tmp_path / "build"))
    rows = [row for item in manifest["materials"] for row in json_rows(Path(item["path"]))]
    assert_source_projection(rows)
    assert (extract["sourceLocale"], extract["targetLocale"], extract["referenceLocales"]) == ("English", "French", ["German"])
    immutable = {row["ContentUid"]: row for item in manifest["materials"] for row in json_rows(Path(item["immutablePath"]))}
    for source in SOURCE:
        assert immutable[source.content_uid]["SourceVersion"] == source.version
    assert immutable["fictional-shared"]["TargetVersion"] == 7


@pytest.mark.parametrize("casings", CASINGS)
def test_project_workflow_prepare_xlsx_preserves_source_and_required_count(tmp_path, casings):
    game = installed_game(tmp_path)
    project = tmp_path / "project.json"
    with backend_scope():
        assert main(["project", "init", "--name", "SYNTHETIC locale regression", "--source", casings[0],
                     "--target", casings[1], "--reference", casings[2], "--game-dir", str(game),
                     "--format", "xlsx", "--workspace", "project-workspace", "--output", str(project)]) == 0
        assert main(["workflow", "prepare", "--project", str(project)]) == 0
    package = tmp_path / "project-workspace/e2e/package"
    manifest = json.loads((package / "translation-package-manifest.json").read_text())
    assert manifest["counts"]["translationRequired"] == 3
    assert (manifest["sourceLocale"], manifest["targetLocale"], manifest["referenceLocales"]) == ("English", "French", ["German"])
    workbook = load_workbook(package / manifest["delivery"]["materials"][0]["path"], read_only=True)
    try:
        values = list(workbook["Translation"].iter_rows(values_only=True))
    finally:
        workbook.close()
    rows = [dict(zip(values[0], row, strict=True)) for row in values[1:]]
    # Empty XLSX cells are represented as None by openpyxl; material source text is still empty.
    for row in rows:
        if row["SourceText"] is None:
            row["SourceText"] = ""
    assert_source_projection(rows)


@pytest.mark.parametrize("consumer", ("build", "package"))
def test_manifest_role_mismatch_cannot_publish_empty_source_success(tmp_path, consumer):
    path, _ = extracted(tmp_path)
    extract = json.loads(path.read_text())
    extract["sourceLocale"] = "english"
    path.write_text(json.dumps(extract), encoding="utf-8")
    output = tmp_path / "must-not-publish"
    with pytest.raises((ValueError, BuildFailure), match="locale|Locale"):
        if consumer == "build":
            run_build(BuildRequest(path, format="jsonl", output=output))
        else:
            prepare_translation_package(extract_manifest=path, project_config_sha256="0" * 64,
                evidence_profile="basic", translation_strategy="standard", material_format="xlsx",
                max_rows=100, output_dir=output)
    assert not output.exists()


@pytest.mark.parametrize("role", ("source", "target", "reference"))
def test_unknown_input_language_rejects_before_extraction(tmp_path, role):
    _, scan = extracted(tmp_path)
    args = {"source": "English", "target": "French", "references": ("German",)}
    args["references" if role == "reference" else role] = ("UnknownLocale",) if role == "reference" else "UnknownLocale"
    output = tmp_path / "unknown"
    with backend_scope(), pytest.raises(ExtractFailure, match="Locale not found"):
        run_extract(ExtractRequest(scan, output=output, **args))
    assert not output.exists()


@pytest.mark.parametrize("condition", ("duplicate_provider", "ambiguous", "unsupported", "duplicate_roles", "ambiguous_entry"))
def test_ambiguous_or_unsupported_input_fails_closed(tmp_path, condition):
    _, scan = extracted(tmp_path)
    data = json.loads(scan.read_text())
    source, target = "english", "french"
    entry = next(r for r in data["locales"] if r["localeId"] == "English")
    if condition == "duplicate_provider":
        copy = deepcopy(entry); copy["localeId"] = "ENGLISH"; data["locales"].append(copy)
    elif condition in {"ambiguous", "unsupported"}:
        entry["status"] = condition
    elif condition == "duplicate_roles":
        target = "ENGLISH"
    else:
        entry["locaEntries"].append({"path": "LOCALIZATION/ENGLISH/ENGLISH.LOCA"})
    scan.write_text(json.dumps(data), encoding="utf-8")
    output = tmp_path / "rejected"
    with backend_scope(), pytest.raises(ExtractFailure):
        run_extract(ExtractRequest(scan, source, target, output=output))
    assert not (output / "extract-manifest.json").exists()


@pytest.mark.parametrize("consumer", ("build", "package"))
@pytest.mark.parametrize("mutation", ("source_none", "wrong_key", "duplicate_key", "text", "version", "count", "normalized_locale", "empty_alignment"))
def test_material_input_inconsistency_rejects_before_publication(tmp_path, consumer, mutation):
    path, _ = extracted(tmp_path)
    extract = json.loads(path.read_text())
    aligned = Path(extract["aligned"])
    rows = json_rows(aligned)
    if mutation == "source_none":
        for row in rows: row["locales"]["English"] = None
    elif mutation in {"wrong_key", "duplicate_key"}:
        for row in rows:
            row["locales"]["english"] = row["locales"]["English"]
            if mutation == "wrong_key": del row["locales"]["English"]
    elif mutation in {"text", "version"}:
        row = next(r for r in rows if r["locales"]["English"] is not None)
        row["locales"]["English"][mutation] = "SYNTHETIC wrong binding" if mutation == "text" else 99
    elif mutation == "count":
        extract["locales"][0]["nodeCount"] += 1
        path.write_text(json.dumps(extract), encoding="utf-8")
    elif mutation == "normalized_locale":
        normalized = Path(extract["locales"][0]["normalized"])
        records = json_rows(normalized); records[0]["localeId"] = "english"
        normalized.write_text("".join(json.dumps(r) + "\n" for r in records), encoding="utf-8")
    else:
        rows = []
    aligned.write_text("".join(json.dumps(r) + "\n" for r in rows), encoding="utf-8")
    output = tmp_path / "must-not-publish"
    with pytest.raises((ValueError, BuildFailure)):
        if consumer == "build":
            run_build(BuildRequest(path, format="jsonl", output=output))
        else:
            prepare_translation_package(extract_manifest=path, project_config_sha256="0" * 64,
                evidence_profile="basic", translation_strategy="standard", material_format="xlsx",
                max_rows=100, output_dir=output)
    assert not output.exists()


@pytest.mark.parametrize("mutation", ("required", "source", "target", "source_version", "target_version", "rows"))
def test_rendered_required_count_and_immutable_projection_are_checked(tmp_path, monkeypatch, mutation):
    path, _ = extracted(tmp_path)
    import bg3loc.builder as module
    render = module.build_material_row

    def faulty_render(*args, **kwargs):
        material, immutable = render(*args, **kwargs)
        if mutation == "required": material["TranslationRequired"] = False
        elif mutation == "source": material["SourceText"] = ""
        elif mutation == "target": material["ExistingTargetText"] = ""
        elif mutation == "source_version": immutable["SourceVersion"] = 99
        elif mutation == "target_version": immutable["TargetVersion"] = 99
        else: material["ContentUid"] = "fictional-duplicate"
        return material, immutable

    monkeypatch.setattr(module, "build_material_row", faulty_render)
    output = tmp_path / "must-not-publish"
    with pytest.raises(BuildFailure, match="locale|Locale"):
        run_build(BuildRequest(path, format="jsonl", output=output))
    assert not output.exists()


@pytest.mark.parametrize("scope,required", [({"fictional-shared"}, 1), ({"fictional-target-only"}, 0),
                                          ({"fictional-reference-only"}, 0)])
def test_explicit_subset_with_legitimate_empty_source_projection(tmp_path, scope, required):
    path, _ = extracted(tmp_path, CASINGS[1])
    package = tmp_path / "subset"
    manifest = prepare_translation_package(extract_manifest=path, project_config_sha256="0" * 64,
        evidence_profile="basic", translation_strategy="standard", material_format="jsonl",
        max_rows=100, output_dir=package, scope_content_uids=scope)
    rows = json_rows(package / manifest["delivery"]["materials"][0]["path"])
    assert {r["ContentUid"] for r in rows} == scope
    assert manifest["counts"]["translationRequired"] == required
    if required == 0:
        assert rows[0]["SourceText"] == "" and rows[0]["TranslationRequired"] is False


def test_empty_explicit_scope_keeps_existing_rejection(tmp_path):
    path, _ = extracted(tmp_path)
    with pytest.raises(ValueError, match="zero localization rows"):
        prepare_translation_package(extract_manifest=path, project_config_sha256="0" * 64,
            evidence_profile="basic", translation_strategy="standard", material_format="jsonl",
            max_rows=100, output_dir=tmp_path / "empty", scope_content_uids=set())


def test_aligned_only_legacy_materials_still_build_when_role_counts_match(tmp_path):
    path, _ = extracted(tmp_path)
    extract = json.loads(path.read_text())
    for entry in extract["locales"]:
        Path(entry["normalized"]).unlink()  # only this test's synthetic sidecars
    result = run_build(BuildRequest(path, format="jsonl", output=tmp_path / "legacy"))
    rows = [r for item in result["materials"] for r in json_rows(Path(item["path"]))]
    assert_source_projection(rows)


def test_cached_broken_manifest_does_not_return_old_empty_source_success(tmp_path, capsys):
    game = installed_game(tmp_path)
    project = tmp_path / "project.json"
    with backend_scope():
        assert main(["project", "init", "--name", "SYNTHETIC cache", "--source", "english", "--target", "french",
                     "--game-dir", str(game), "--format", "jsonl", "--workspace", "work", "--output", str(project)]) == 0
        assert main(["workflow", "prepare", "--project", str(project)]) == 0
        extract = tmp_path / "work/e2e/extract/extract-manifest.json"
        data = json.loads(extract.read_text()); data["sourceLocale"] = "english"
        extract.write_text(json.dumps(data), encoding="utf-8")
        # Model a pre-fix reusable extract checkpoint, including its current file seal.
        state_path = tmp_path / "work/e2e/workflow-state.json"
        state = json.loads(state_path.read_text())
        from bg3loc.e2e.package import sha256_file
        state["stages"]["extract"]["sha256"] = sha256_file(extract)
        state_path.write_text(json.dumps(state), encoding="utf-8")
        package = tmp_path / "work/e2e/package"
        before = {p.relative_to(package): p.read_bytes() for p in package.rglob("*") if p.is_file()}
        assert main(["workflow", "prepare", "--project", str(project)]) == 73
        assert "locale roles are not canonical" in capsys.readouterr().err
        assert {p.relative_to(package): p.read_bytes() for p in package.rglob("*") if p.is_file()} == before


def test_full_workflow_forwards_resolved_canonical_roles_to_research(tmp_path):
    game = installed_game(tmp_path)
    project = tmp_path / "project.json"
    observed = []

    def scan(request):
        observed.append((request.source, request.target, request.references))
        request.output.parent.mkdir(parents=True, exist_ok=True)
        request.output.write_text("{}", encoding="utf-8")
        return 0

    def mappings(request):
        observed.append((request.source, request.target, request.references))
        request.output_dir.mkdir(parents=True, exist_ok=True)
        (request.output_dir / "research-summary.json").write_text("{}", encoding="utf-8")
        (request.output_dir / "research-mappings.jsonl").write_text("", encoding="utf-8")
        return 0

    with backend_scope(), patch("bg3loc.e2e.prepare.run_research_scan_request", side_effect=scan), \
         patch("bg3loc.e2e.prepare.run_research_map_request", side_effect=mappings), \
         patch("bg3loc.e2e.prepare.research_evidence_sidecars", return_value={}):
        assert main(["project", "init", "--name", "SYNTHETIC full workflow", "--source", "english", "--target", "french",
                     "--reference", "german", "--game-dir", str(game), "--profile", "full", "--format", "jsonl",
                     "--workspace", "work", "--output", str(project)]) == 0
        assert main(["workflow", "prepare", "--project", str(project)]) == 0
    assert observed == [("English", "French", ("German",))] * 2
    assert json.loads(project.read_text())["locales"]["source"] == "english"  # user config stays intact
