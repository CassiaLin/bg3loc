from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


SCRIPT_PATH = Path(__file__).resolve().parents[1] / "scripts" / "extract-research-baseline.py"
SPEC = importlib.util.spec_from_file_location("extract_research_baseline", SCRIPT_PATH)
assert SPEC is not None and SPEC.loader is not None
EXTRACTOR = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(EXTRACTOR)


SHARED_SPEAKER_HEADERS = [
    "ReviewItemId",
    "ContentUid",
    "ConflictGroup",
    "EnglishText",
    "CurrentTraditionalChineseText",
    "SpeakerIdentity",
    "TranslationConstraint",
    "ReviewDecision",
    "ReviewerTranslation",
    "ReviewerNote",
]


class ResearchBaselineHeaderTests(unittest.TestCase):
    def test_normal_header_in_first_row_is_unchanged(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Normal"
        sheet.append(["ContentUid", "Value"])
        sheet.append(["h001", "first"])
        sheet.append(["h002", "second"])

        rows = list(EXTRACTOR._rows_as_dicts(sheet))

        self.assertEqual(
            rows,
            [
                (2, {"ContentUid": "h001", "Value": "first"}),
                (3, {"ContentUid": "h002", "Value": "second"}),
            ],
        )

    def test_notice_before_header_is_detected(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Shifted"
        sheet.append(["Read this notice before editing."])
        sheet.append(["ContentUid", "Value"])
        sheet.append(["h001", "first"])

        rows = list(EXTRACTOR._rows_as_dicts(sheet))

        self.assertEqual(rows, [(3, {"ContentUid": "h001", "Value": "first"})])

    def test_missing_header_fails_closed(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "MissingHeader"
        for row_number in range(1, 11):
            sheet.append([f"notice {row_number}"])

        with self.assertRaisesRegex(
            ValueError,
            r"Sheet 'MissingHeader'.*rows 1-10.*ContentUid",
        ):
            list(EXTRACTOR._rows_as_dicts(sheet))

    def test_shared_speaker_shape_extracts_twelve_bark_reviews(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "SharedSpeaker"
        sheet.append(["IMPORTANT: Translation must fit all shared-set members."])
        sheet.append(SHARED_SPEAKER_HEADERS)
        for ordinal in range(1, 13):
            sheet.append(
                [
                    f"BARK-SHARED-{ordinal:03d}",
                    f"h{ordinal:032x}",
                    f"CONFLICT-{ordinal:03d}",
                    f"English {ordinal}",
                    "",
                    "Shared speaker set; no single member selected",
                    "MustFitAllSharedSpeakers",
                    "Approved",
                    f"Translation {ordinal}",
                    f"Note {ordinal}",
                ]
            )

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "synthetic-bark.xlsx"
            workbook.save(path)
            by_uid = {}
            EXTRACTOR._ingest_bark(path, by_uid)

        reviews = [review for record in by_uid.values() for review in record.get("barkReview", [])]
        self.assertEqual(len(reviews), 12)
        self.assertTrue(all(review["provenance"]["sheet"] == "SharedSpeaker" for review in reviews))
        self.assertEqual(reviews[0]["provenance"]["row"], 3)
        self.assertEqual(reviews[-1]["provenance"]["row"], 14)
        self.assertEqual(reviews[0]["TranslationConstraint"], "MustFitAllSharedSpeakers")
        self.assertEqual(reviews[0]["ReviewDecision"], "Approved")

    def test_production_shape_notice_blank_header_is_detected(self) -> None:
        # The preserved production workbook was inspected before this fix:
        # row 1 is a notice, row 2 is blank, row 3 is the actual header, followed
        # by 12 SharedSpeaker data rows. This synthetic shape records that evidence.
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "SharedSpeaker"
        sheet.append(["IMPORTANT: notice"])
        sheet.append([])
        sheet.append(SHARED_SPEAKER_HEADERS)
        sheet.append(["BARK-SHARED-001", "h001"])

        rows = list(EXTRACTOR._rows_as_dicts(sheet))

        self.assertEqual(rows[0][0], 4)
        self.assertEqual(rows[0][1]["ContentUid"], "h001")


if __name__ == "__main__":
    unittest.main()
