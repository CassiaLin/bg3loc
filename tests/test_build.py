from __future__ import annotations

import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from bg3loc.builder import BuildFailure, BuildRequest, extract_protected_tokens, run_build


class BuildV1Tests(unittest.TestCase):
    def _make_extract_fixture(self, root: Path) -> Path:
        aligned = root / "extract" / "aligned" / "source-target.jsonl"
        aligned.parent.mkdir(parents=True)
        rows = [
            {
                "contentUid": "h001",
                "locales": {
                    "English": {"text": "Hello <LSTag Type=\"Spell\" Tooltip=\"Source_Action\">{Player}</LSTag> [Narrative]", "version": 1, "presence": True},
                    "ChineseTraditional": {"text": "你好 <LSTag Tooltip=\"Localized_Action\">{Player}</LSTag> [本地化]", "version": 1, "presence": True},
                    "Russian": {"text": "Привет <LSTag Tooltip=\"Localized_Action\">{Player}</LSTag>", "version": 1, "presence": True},
                },
            },
            {
                "contentUid": "h002",
                "locales": {
                    "English": {"text": "Source only", "version": 2, "presence": True},
                    "ChineseTraditional": None,
                    "Russian": {"text": "Источник", "version": 2, "presence": True},
                },
            },
            {
                "contentUid": "h003",
                "locales": {
                    "English": None,
                    "ChineseTraditional": {"text": "目標語言限定", "version": 3, "presence": True},
                    "Russian": None,
                },
            },
        ]
        aligned.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")

        validation = root / "extract" / "validation" / "roundtrip.json"
        validation.parent.mkdir(parents=True)
        validation.write_text("{}", encoding="utf-8")

        manifest = {
            "schemaVersion": "1.0",
            "scanManifest": str(root / "scan-manifest.json"),
            "sourceLocale": "English",
            "targetLocale": "ChineseTraditional",
            "referenceLocales": ["Russian"],
            "backend": {"id": "fake"},
            "locales": [
                {
                    "localeId": "English",
                    "packageFile": "English.pak",
                    "locaEntry": "english.loca",
                    "sourceLoca": "English/source.loca",
                    "sourceXml": "English/source.xml",
                    "normalized": "English.jsonl",
                    "nodeCount": 2,
                },
                {
                    "localeId": "ChineseTraditional",
                    "packageFile": "ChineseTraditional.pak",
                    "locaEntry": "chinesetraditional.loca",
                    "sourceLoca": "ChineseTraditional/source.loca",
                    "sourceXml": "ChineseTraditional/source.xml",
                    "normalized": "ChineseTraditional.jsonl",
                    "nodeCount": 2,
                },
                {
                    "localeId": "Russian",
                    "packageFile": "Russian.pak",
                    "locaEntry": "russian.loca",
                    "sourceLoca": "Russian/source.loca",
                    "sourceXml": "Russian/source.xml",
                    "normalized": "Russian.jsonl",
                    "nodeCount": 2,
                },
            ],
            "aligned": str(aligned),
            "roundtripValidation": str(validation),
        }
        manifest_path = root / "extract" / "extract-manifest.json"
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        return manifest_path

    def test_basic_csv_includes_existing_target_reference_and_splits(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            extract = self._make_extract_fixture(root)
            output = root / "build"
            manifest = run_build(BuildRequest(extract_manifest=extract, mode="basic", format="csv", max_rows=2, output=output))

            self.assertEqual(len(manifest["materials"]), 2)
            first = Path(manifest["materials"][0]["path"])
            with first.open("r", encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.DictReader(stream))
            self.assertEqual(len(rows), 2)
            self.assertIn("ExistingTargetText", rows[0])
            self.assertIn("ReferenceText[Russian]", rows[0])
            self.assertEqual(rows[0]["TranslatedTargetText"], "")
            self.assertEqual(rows[1]["PresenceStatus"], "source-only")

            immutable = Path(manifest["materials"][0]["immutablePath"])
            self.assertTrue(immutable.is_file())
            self.assertEqual(len(manifest["materials"][0]["sha256"]), 64)

    def test_blind_first_xlsx_physically_hides_target_and_reference(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            extract = self._make_extract_fixture(root)
            output = root / "build"
            manifest = run_build(BuildRequest(extract_manifest=extract, mode="blind-first", format="xlsx", output=output))

            material = Path(manifest["materials"][0]["path"])
            workbook = load_workbook(material, read_only=True)
            sheet = workbook["Translation"]
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            self.assertNotIn("ExistingTargetText", headers)
            self.assertNotIn("ReferenceText[Russian]", headers)
            self.assertIn("IndependentTargetText", headers)
            self.assertIn("ProtectedTokens", headers)

            data = list(sheet.iter_rows(min_row=2, max_row=2, values_only=True))[0]
            row = dict(zip(headers, data, strict=True))
            self.assertEqual(row["SourceText"], "Hello <LSTag Type=\"Spell\" Tooltip=\"Source_Action\">{Player}</LSTag> [Narrative]")
            self.assertIsNone(row["IndependentTargetText"])
            self.assertNotIn("你好", json.dumps(row, ensure_ascii=False))
            self.assertNotIn("Привет", json.dumps(row, ensure_ascii=False))
            workbook.close()

            validation = json.loads((output / "validation" / "build-validation.json").read_text(encoding="utf-8"))
            self.assertTrue(validation["blindFirstPhysicalSeparation"])
            self.assertFalse(validation["translationCandidateGeneration"])

    def test_context_mode_requires_context(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            extract = self._make_extract_fixture(root)
            with self.assertRaises(BuildFailure) as ctx:
                run_build(BuildRequest(extract_manifest=extract, mode="context", format="jsonl", output=root / "build"))
            self.assertEqual(ctx.exception.exit_code, 22)

    def test_protected_token_extraction_keeps_only_exact_runtime_lexemes_and_multiplicity(self) -> None:
        tokens = extract_protected_tokens("<b>{Player}</b> {Player} %s [FLAG_Test]")
        self.assertEqual(tokens, ["{Player}", "{Player}", "%s"])


if __name__ == "__main__":
    unittest.main()
