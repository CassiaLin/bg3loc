from __future__ import annotations

import csv
import glob
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from bg3loc.io import read_json, write_json
from bg3loc.protected_syntax import (
    extract_protected_tokens,
    has_valid_lstag_structure,
    validate_protected_syntax,
)
from bg3loc.schema import SchemaStore


@dataclass(frozen=True, slots=True)
class ValidateRequest:
    build_manifest: Path
    input_spec: str
    strict: bool = True
    output: Path = Path("workspace/validate")


class ValidateFailure(RuntimeError):
    def __init__(self, exit_code: int, key: str, message: str) -> None:
        super().__init__(message)
        self.exit_code = exit_code
        self.key = key
        self.message = message


@dataclass(slots=True)
class Finding:
    code: str
    severity: str
    message: str
    content_uid: str | None = None
    file: str | None = None
    row: int | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "code": self.code,
            "severity": self.severity,
            "message": self.message,
            "contentUid": self.content_uid,
            "file": self.file,
            "row": self.row,
        }


_INTENTIONAL_EMPTY = {"intentionally-empty", "empty", "not-applicable", "skip"}
_CONTEXT_STATUSES = {"needs-more-context", "needs_more_context", "more-context"}
_ESCALATE_STATUSES = {"escalate", "escalated"}


def run_validate(request: ValidateRequest) -> dict[str, Any]:
    try:
        build = read_json(request.build_manifest)
        SchemaStore().validate("build-manifest.schema.json", build)
    except Exception as exc:
        raise ValidateFailure(30, "BUILD_MANIFEST_INVALID", str(exc)) from exc

    try:
        workbook_schema_path = resolve_build_reference(
            str(build.get("workbookSchema", "")), request.build_manifest.parent, fallback_subdir="schemas"
        )
        workbook_schema = read_json(workbook_schema_path)
    except Exception as exc:
        raise ValidateFailure(30, "BUILD_MANIFEST_INVALID", f"Unable to load workbook schema: {exc}") from exc

    try:
        baseline = load_immutable_baseline(build, request.build_manifest.parent)
    except Exception as exc:
        raise ValidateFailure(30, "BUILD_MANIFEST_INVALID", f"Unable to load immutable baseline: {exc}") from exc

    input_files = expand_input_spec(request.input_spec)
    if not input_files:
        raise ValidateFailure(31, "RETURN_MATERIAL_NOT_FOUND", f"No return material matched: {request.input_spec}")

    returned_rows: list[dict[str, Any]] = []
    try:
        for path in input_files:
            returned_rows.extend(read_return_material(path))
    except Exception as exc:
        raise ValidateFailure(32, "RETURN_FORMAT_INVALID", str(exc)) from exc

    visible_fields = [str(item) for item in workbook_schema.get("visibleFields", [])]
    editable_fields = {str(item) for item in workbook_schema.get("editableFields", [])}
    if "ContentUid" not in visible_fields:
        raise ValidateFailure(30, "BUILD_MANIFEST_INVALID", "workbook schema does not expose ContentUid")

    mode = str(build["mode"])
    target_locale = str(build["targetLocale"])
    findings: list[Finding] = []
    seen_return_uids: dict[str, int] = {}
    returned_uid_set: set[str] = set()
    evaluated_required: set[str] = set()
    accepted_records: list[dict[str, Any]] = []
    rejected_records: list[dict[str, Any]] = []

    for index, row in enumerate(returned_rows):
        row_no = int(row.pop("__row__", index + 1))
        file_name = str(row.pop("__file__", ""))
        row_errors: list[Finding] = []
        row_warnings: list[Finding] = []
        content_uid = normalize_text(row.get("ContentUid"))

        if not content_uid or content_uid not in baseline:
            row_errors.append(Finding("VAL001", "error", "UNKNOWN_CONTENT_UID", content_uid or None, file_name, row_no))
        elif content_uid in seen_return_uids:
            row_errors.append(Finding("VAL002", "error", "DUPLICATE_CONTENT_UID", content_uid, file_name, row_no))
        else:
            seen_return_uids[content_uid] = index
            returned_uid_set.add(content_uid)
            immutable = baseline[content_uid]
            row_errors.extend(validate_immutable_fields(row, immutable, visible_fields, editable_fields, file_name, row_no))

            status = normalize_text(row.get("TranslationStatus")).casefold()
            if bool(immutable.get("TranslationRequired")) and status:
                evaluated_required.add(content_uid)
            if status in _CONTEXT_STATUSES:
                row_warnings.append(Finding("VAL102", "warning", "STATUS_NEEDS_MORE_CONTEXT", content_uid, file_name, row_no))
            if status in _ESCALATE_STATUSES:
                row_warnings.append(Finding("VAL103", "warning", "STATUS_ESCALATE", content_uid, file_name, row_no))

            candidate_field = "IndependentTargetText" if mode == "blind-first" else "TranslatedTargetText"
            candidate = normalize_text(row.get(candidate_field))
            if candidate:
                protected_tokens = immutable.get("ProtectedTokens", [])
                if not isinstance(protected_tokens, list) or not all(isinstance(item, str) for item in protected_tokens):
                    row_errors.append(Finding("VAL008", "error", "MARKUP_INVALID: invalid protected syntax baseline", content_uid, file_name, row_no))
                else:
                    row_errors.extend(validate_protected_content(protected_tokens, candidate, content_uid, file_name, row_no))
                    if (
                        immutable.get("PresenceStatus") == "source-only"
                        and has_valid_lstag_structure(str(immutable.get("SourceText", "")))
                        and not has_valid_lstag_structure(candidate)
                        and not any(item.code == "VAL008" for item in row_errors)
                    ):
                        row_errors.append(Finding(
                            "VAL008",
                            "error",
                            "MARKUP_INVALID: source-only LSTag structure removed",
                            content_uid,
                            file_name,
                            row_no,
                        ))
            elif bool(immutable.get("TranslationRequired")) and status and status not in _INTENTIONAL_EMPTY:
                row_errors.append(Finding("VAL009", "error", "REQUIRED_TARGET_EMPTY", content_uid, file_name, row_no))

            if not row_errors and should_accept_candidate(candidate, status, immutable, mode):
                accepted_records.append({
                    "contentUid": content_uid,
                    "localeId": target_locale,
                    "text": candidate,
                    "version": immutable.get("TargetVersion") if immutable.get("TargetVersion") is not None else immutable.get("SourceVersion"),
                    "translationStatus": normalize_text(row.get("TranslationStatus")),
                    "translatorNotes": normalize_text(row.get("TranslatorNotes")),
                    "translatorName": normalize_text(row.get("TranslatorName")),
                })

        if row_errors or row_warnings:
            findings.extend(row_errors + row_warnings)
        if row_errors:
            rejected_records.append({
                "contentUid": content_uid or None,
                "file": file_name,
                "row": row_no,
                "errors": [item.as_dict() for item in row_errors],
                "returned": {key: value for key, value in row.items()},
            })

    missing_return_row_count = len(set(baseline) - returned_uid_set)

    required_uids = {uid for uid, item in baseline.items() if bool(item.get("TranslationRequired"))}
    coverage: float | None = None
    if mode == "blind-first":
        coverage = 1.0 if not required_uids else len(evaluated_required & required_uids) / len(required_uids)
        if coverage < 1.0:
            for content_uid in sorted(required_uids - evaluated_required):
                findings.append(Finding("VAL010", "error", "BLIND_FIRST_COVERAGE_INCOMPLETE", content_uid))

    error_count = sum(1 for item in findings if item.severity == "error")
    warning_count = sum(1 for item in findings if item.severity == "warning")
    if error_count == 0:
        status = "pass"
    elif request.strict or not accepted_records:
        status = "fail"
    else:
        status = "partial"

    request.output.mkdir(parents=True, exist_ok=True)
    accepted_path = request.output / "accepted" / "normalized-target.jsonl"
    rejected_path = request.output / "rejected" / "rejected-rows.jsonl"
    summary_path = request.output / "reports" / "validation-summary.json"
    errors_path = request.output / "reports" / "validation-errors.csv"
    try:
        write_jsonl(accepted_path, accepted_records)
        write_jsonl(rejected_path, rejected_records)
        write_findings_csv(errors_path, findings)
        write_json(summary_path, {
            "schemaVersion": "1.0",
            "status": status,
            "strict": request.strict,
            "inputFiles": [str(path) for path in input_files],
            "returnedRowCount": len(returned_rows),
            "baselineRowCount": len(baseline),
            "acceptedRecordCount": len(accepted_records),
            "rejectedRowCount": len(rejected_records),
            "errorCount": error_count,
            "warningCount": warning_count,
            "missingReturnRowCount": missing_return_row_count,
            "independentEvaluationCoverage": coverage,
        })

        manifest = {
            "schemaVersion": "1.0",
            "buildManifest": str(request.build_manifest),
            "strict": request.strict,
            "accepted": {"path": str(accepted_path), "recordCount": len(accepted_records)},
            "rejected": {"path": str(rejected_path), "recordCount": len(rejected_records)},
            "reports": {"summary": str(summary_path), "errors": str(errors_path)},
            "summary": {
                "status": status,
                "errorCount": error_count,
                "warningCount": warning_count,
                "missingReturnRowCount": missing_return_row_count,
                "independentEvaluationCoverage": coverage,
            },
        }
        SchemaStore().validate("validate-manifest.schema.json", manifest)
        write_json(request.output / "validate-manifest.json", manifest)
    except Exception as exc:
        raise ValidateFailure(37, "VALIDATION_OUTPUT_FAILED", str(exc)) from exc

    if request.strict and error_count:
        code, key = exit_boundary_for_findings(findings)
        raise ValidateFailure(code, key, f"Validation failed with {error_count} error(s); outputs were written to {request.output}")
    return manifest


def load_immutable_baseline(build: dict[str, Any], manifest_dir: Path) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for material in build.get("materials", []):
        raw = material.get("immutablePath")
        if not raw:
            raise ValueError("build material is missing immutablePath")
        path = resolve_build_reference(str(raw), manifest_dir, fallback_subdir="immutable")
        for row in read_jsonl(path):
            uid = normalize_text(row.get("ContentUid"))
            if not uid:
                raise ValueError(f"immutable row missing ContentUid: {path}")
            if uid in result:
                raise ValueError(f"duplicate immutable ContentUid: {uid}")
            result[uid] = row
    return result


def validate_immutable_fields(
    returned: dict[str, Any],
    immutable: dict[str, Any],
    visible_fields: list[str],
    editable_fields: set[str],
    file_name: str,
    row_no: int,
) -> list[Finding]:
    uid = str(immutable["ContentUid"])
    findings: list[Finding] = []

    def mismatch(field: str, expected: Any, code: str, label: str) -> None:
        if field not in returned:
            return
        if canonical_cell(returned.get(field)) != canonical_cell(expected):
            findings.append(Finding(code, "error", label, uid, file_name, row_no))

    mismatch("SourceText", immutable.get("SourceText", ""), "VAL003", "SOURCE_TEXT_MODIFIED")
    mismatch("SourceVersion", immutable.get("SourceVersion"), "VAL004", "SOURCE_VERSION_MODIFIED")
    mismatch("SourceLocale", immutable.get("SourceLocale"), "VAL005", "LOCALE_ROLE_MODIFIED")
    mismatch("TargetLocale", immutable.get("TargetLocale"), "VAL005", "LOCALE_ROLE_MODIFIED")
    mismatch("PresenceStatus", immutable.get("PresenceStatus"), "VAL011", "UNEXPECTED_EDITED_FIELD")
    mismatch("TranslationRequired", immutable.get("TranslationRequired"), "VAL011", "UNEXPECTED_EDITED_FIELD")
    mismatch("ExistingTargetText", immutable.get("ExistingTargetText", ""), "VAL011", "UNEXPECTED_EDITED_FIELD")
    mismatch("TargetVersion", immutable.get("TargetVersion"), "VAL011", "UNEXPECTED_EDITED_FIELD")
    mismatch("ProtectedTokens", json.dumps(immutable.get("ProtectedTokens", []), ensure_ascii=False, separators=(",", ":")), "VAL011", "UNEXPECTED_EDITED_FIELD")

    references = immutable.get("References", {})
    if isinstance(references, dict):
        for locale_id, value in references.items():
            text = "" if value is None else str(value.get("text", ""))
            mismatch(f"ReferenceText[{locale_id}]", text, "VAL011", "UNEXPECTED_EDITED_FIELD")

    if "Context" in visible_fields:
        mismatch("Context", immutable.get("Context"), "VAL011", "UNEXPECTED_EDITED_FIELD")

    allowed = set(visible_fields)
    for field, value in returned.items():
        if isinstance(field, str) and (field.startswith("__") or field in allowed):
            continue
        if canonical_cell(value) != "":
            findings.append(Finding("VAL011", "error", f"UNEXPECTED_EDITED_FIELD: {field}", uid, file_name, row_no))
    return findings


def validate_protected_content(required_tokens: list[str], candidate: str, uid: str, file_name: str, row_no: int) -> list[Finding]:
    findings: list[Finding] = []
    for issue in validate_protected_syntax(required_tokens, candidate):
        if issue.kind == "missing":
            findings.append(Finding("VAL006", "error", f"PROTECTED_TOKEN_MISSING: {list(issue.details[:10])}", uid, file_name, row_no))
        elif issue.kind == "added":
            findings.append(Finding("VAL007", "error", f"PROTECTED_TOKEN_ADDED: {list(issue.details[:10])}", uid, file_name, row_no))
        else:
            findings.append(Finding("VAL008", "error", f"MARKUP_INVALID: {list(issue.details[:10])}", uid, file_name, row_no))
    return findings


def token_sequence(text: str) -> list[str]:
    return extract_protected_tokens(text)


def should_accept_candidate(candidate: str, status: str, immutable: dict[str, Any], mode: str) -> bool:
    if candidate:
        return True
    return status in _INTENTIONAL_EMPTY and bool(immutable.get("TranslationRequired"))


def exit_boundary_for_findings(findings: list[Finding]) -> tuple[int, str]:
    codes = {item.code for item in findings if item.severity == "error"}
    if codes & {"VAL001", "VAL002"}:
        return 33, "IDENTITY_VALIDATION_FAILED"
    if codes & {"VAL003", "VAL004", "VAL005", "VAL011"}:
        return 34, "IMMUTABLE_FIELD_VALIDATION_FAILED"
    if codes & {"VAL006", "VAL007", "VAL008"}:
        return 35, "PROTECTED_TOKEN_VALIDATION_FAILED"
    return 36, "COVERAGE_VALIDATION_FAILED"


def expand_input_spec(spec: str) -> list[Path]:
    direct = Path(spec).expanduser()
    if direct.is_file():
        return [direct.resolve()]
    if direct.is_dir():
        files: list[Path] = []
        for suffix in ("*.xlsx", "*.csv", "*.jsonl"):
            files.extend(direct.glob(suffix))
        return sorted({path.resolve() for path in files}, key=lambda path: str(path).casefold())
    matches = [Path(item).resolve() for item in glob.glob(spec, recursive=True) if Path(item).is_file()]
    return sorted(set(matches), key=lambda path: str(path).casefold())


def read_return_material(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.casefold()
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=False)
        if "Translation" not in workbook.sheetnames:
            raise ValueError(f"XLSX missing Translation sheet: {path}")
        sheet = workbook["Translation"]
        iterator = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(iterator)
        except StopIteration:
            return []
        headers = [normalize_text(value) for value in raw_headers]
        if not headers or "ContentUid" not in headers:
            raise ValueError(f"XLSX missing ContentUid header: {path}")
        rows: list[dict[str, Any]] = []
        for row_no, values in enumerate(iterator, start=2):
            if all(value is None or value == "" for value in values):
                continue
            row = {header: value for header, value in zip(headers, values, strict=False) if header}
            row["__file__"] = str(path)
            row["__row__"] = row_no
            rows.append(row)
        return rows
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            if not reader.fieldnames or "ContentUid" not in reader.fieldnames:
                raise ValueError(f"CSV missing ContentUid header: {path}")
            rows = []
            for row_no, row in enumerate(reader, start=2):
                if not any(value not in (None, "") for value in row.values()):
                    continue
                item = dict(row)
                item["__file__"] = str(path)
                item["__row__"] = row_no
                rows.append(item)
            return rows
    if suffix == ".jsonl":
        rows = read_jsonl(path)
        for row_no, row in enumerate(rows, start=1):
            row["__file__"] = str(path)
            row["__row__"] = row_no
        return rows
    raise ValueError(f"Unsupported return format: {path}")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as stream:
        for line_no, line in enumerate(stream, start=1):
            if not line.strip():
                continue
            value = json.loads(line)
            if not isinstance(value, dict):
                raise ValueError(f"JSONL line {line_no} is not an object: {path}")
            rows.append(value)
    return rows


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with tmp.open("w", encoding="utf-8", newline="\n") as stream:
        for row in rows:
            stream.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n")
    tmp.replace(path)


def write_findings_csv(path: Path, findings: list[Finding]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with tmp.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["code", "severity", "contentUid", "file", "row", "message"])
        writer.writeheader()
        for item in findings:
            writer.writerow({
                "code": item.code,
                "severity": item.severity,
                "contentUid": item.content_uid or "",
                "file": item.file or "",
                "row": "" if item.row is None else item.row,
                "message": item.message,
            })
    tmp.replace(path)


def resolve_build_reference(raw: str, manifest_dir: Path, fallback_subdir: str | None = None) -> Path:
    candidate = Path(raw).expanduser()
    if candidate.is_file():
        return candidate.resolve()
    if not candidate.is_absolute():
        relative = manifest_dir / candidate
        if relative.is_file():
            return relative.resolve()
    if fallback_subdir:
        fallback = manifest_dir / fallback_subdir / candidate.name
        if fallback.is_file():
            return fallback.resolve()
    raise FileNotFoundError(raw)


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def canonical_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, str):
        stripped = value.strip()
        lowered = stripped.casefold()
        if lowered in {"true", "false"}:
            return lowered
        if stripped.startswith(("{", "[")):
            try:
                parsed = json.loads(stripped)
            except json.JSONDecodeError:
                return value
            return canonical_cell(parsed)
        return value
    return str(value)
