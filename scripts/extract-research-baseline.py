from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


CONTEXT_COLUMNS = [
    "ContentUid",
    "ConflictGroup",
    "ContextSource",
    "ContextType",
    "ContextSummary",
    "EvidenceStrength",
    "TranslationConstraint",
    "PakName",
    "InternalPath",
    "DialogUUID",
    "NodeUUID",
]

BARK_EVIDENCE_COLUMNS = [
    "ReviewItemId",
    "ContentUid",
    "ConflictGroup",
    "PakName",
    "InternalPath",
    "DialogUUID",
    "NodeUUID",
    "ContainerPurposeEvidence",
    "SpeakerEvidence",
    "ContextStrength",
    "ReviewReadiness",
    "EvidenceNotes",
    "SourceReport",
]

QUEST_REVIEW_COLUMNS = [
    "ReviewRowId",
    "ContentUid",
    "ConflictGroup",
    "EnglishVersion",
    "EnglishText",
    "QuestJournalFieldRole",
    "EvidenceStrength",
    "PakName",
    "InternalPath",
    "Module",
    "RecordType",
    "RecordUUID／RecordId",
    "NodeType",
    "NodeUUID／NodeId",
    "FieldName",
    "FieldPath",
    "StructuralOwner",
    "PhysicalOccurrenceCount",
    "OccurrenceClassification",
    "TraditionalChineseAvailability",
    "LocalizationActionStatus",
    "TranslationConstraint",
    "ReviewDecision",
    "ReviewerTranslation",
    "ReviewerNote",
    "ReviewerName",
    "ReviewedAt",
]

QUEST_AVAILABILITY_COLUMNS = [
    "ContentUid",
    "EnglishXmlExactHitCount",
    "EnglishVersion",
    "TraditionalChineseXmlExactHitCount",
    "AvailabilityStatus",
    "LocalizationActionStatus",
    "AdditionAuthorized",
    "TranslationApproved",
]

REVIEW_COLUMNS = [
    "ReviewItemId",
    "ContentUid",
    "EnglishVersion",
    "EnglishText",
    "EnglishTextStatus",
    "CandidateClassification",
    "ContextAvailability",
    "ContextType",
    "ContextSummary",
    "TranslationConstraint",
    "ReviewDecision",
    "ReviewerTranslation",
    "ReviewerNote",
    "ReviewerName",
    "ReviewedAt",
]


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Extract preserved BG3 research evidence from translation workbooks")
    p.add_argument("--full", type=Path, help="bg3-english-only-full-translation-review.xlsx")
    p.add_argument("--bark", type=Path, help="phase4w-bark-human-review.xlsx")
    p.add_argument("--quest", type=Path, help="phase4ad-quest-human-review.xlsx")
    p.add_argument("--four-language", type=Path, dest="four_language", help="bg3-four-language-translator-table-filled.xlsx")
    p.add_argument("--output", type=Path, default=Path("workspace/research-baseline"))
    return p.parse_args()


def main() -> int:
    args = parse_args()
    sources = [path for path in (args.full, args.bark, args.quest, args.four_language) if path]
    if not sources:
        raise SystemExit("Provide at least one workbook: --full, --bark, --quest, or --four-language")

    out = args.output
    out.mkdir(parents=True, exist_ok=True)

    by_uid: dict[str, dict[str, Any]] = {}
    source_manifest: list[dict[str, Any]] = []

    if args.full:
        _ingest_full(args.full, by_uid)
        source_manifest.append(_source_entry(args.full))
    if args.bark:
        _ingest_bark(args.bark, by_uid)
        source_manifest.append(_source_entry(args.bark))
    if args.quest:
        _ingest_quest(args.quest, by_uid)
        source_manifest.append(_source_entry(args.quest))
    if args.four_language:
        _ingest_four_language(args.four_language, by_uid)
        source_manifest.append(_source_entry(args.four_language))

    baseline_path = out / "research-baseline.jsonl"
    with baseline_path.open("w", encoding="utf-8", newline="\n") as fh:
        for uid in sorted(by_uid):
            fh.write(json.dumps(by_uid[uid], ensure_ascii=False, separators=(",", ":")) + "\n")

    summary = _summarize(by_uid)
    summary["schemaVersion"] = "1.1"
    summary["sources"] = source_manifest
    summary["baseline"] = str(baseline_path)
    (out / "research-baseline-summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return 0


def _ingest_full(path: Path, by_uid: dict[str, dict[str, Any]]) -> None:
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if "ContextEvidence" in wb.sheetnames:
            for row_number, row in _rows_as_dicts(wb["ContextEvidence"]):
                uid = _text(row.get("ContentUid"))
                if not uid:
                    continue
                record = _record(by_uid, uid)
                record.setdefault("contextEvidence", []).append(
                    _project(row, CONTEXT_COLUMNS, provenance=_prov(path, "ContextEvidence", row_number))
                )
        if "TranslationQueue" in wb.sheetnames:
            for row_number, row in _rows_as_dicts(wb["TranslationQueue"]):
                uid = _text(row.get("ContentUid"))
                if not uid:
                    continue
                record = _record(by_uid, uid)
                record.setdefault("reviewEvidence", []).append(
                    _project(row, REVIEW_COLUMNS, provenance=_prov(path, "TranslationQueue", row_number))
                )
    finally:
        wb.close()


def _ingest_bark(path: Path, by_uid: dict[str, dict[str, Any]]) -> None:
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if "EvidenceReference" in wb.sheetnames:
            for row_number, row in _rows_as_dicts(wb["EvidenceReference"]):
                uid = _text(row.get("ContentUid"))
                if not uid:
                    continue
                record = _record(by_uid, uid)
                record.setdefault("barkEvidence", []).append(
                    _project(row, BARK_EVIDENCE_COLUMNS, provenance=_prov(path, "EvidenceReference", row_number))
                )
        for sheet_name in ("SingleSpeaker", "SharedSpeaker"):
            if sheet_name not in wb.sheetnames:
                continue
            for row_number, row in _rows_as_dicts(wb[sheet_name]):
                uid = _text(row.get("ContentUid"))
                if not uid:
                    continue
                record = _record(by_uid, uid)
                payload = {key: _json_cell(value) for key, value in row.items() if key and value not in (None, "")}
                payload["provenance"] = _prov(path, sheet_name, row_number)
                record.setdefault("barkReview", []).append(payload)
        if "SharedSetMembers" in wb.sheetnames:
            shared_rows = []
            for row_number, row in _rows_as_dicts(
                wb["SharedSetMembers"], required_headers={"SharedSetId", "MemberUUID"}
            ):
                payload = {key: _json_cell(value) for key, value in row.items() if key and value not in (None, "")}
                payload["provenance"] = _prov(path, "SharedSetMembers", row_number)
                shared_rows.append(payload)
            if shared_rows:
                by_uid.setdefault("__shared_sets__", {"contentUid": "__shared_sets__"})["sharedSetMembers"] = shared_rows
    finally:
        wb.close()


def _ingest_quest(path: Path, by_uid: dict[str, dict[str, Any]]) -> None:
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if "QuestReview" in wb.sheetnames:
            for row_number, row in _rows_as_dicts(wb["QuestReview"]):
                uid = _text(row.get("ContentUid"))
                if not uid:
                    continue
                record = _record(by_uid, uid)
                record.setdefault("questReview", []).append(
                    _project(row, QUEST_REVIEW_COLUMNS, provenance=_prov(path, "QuestReview", row_number))
                )
        if "LocalizationAvailability" in wb.sheetnames:
            for row_number, row in _rows_as_dicts(wb["LocalizationAvailability"]):
                uid = _text(row.get("ContentUid"))
                if not uid:
                    continue
                record = _record(by_uid, uid)
                record.setdefault("questLocalizationAvailability", []).append(
                    _project(
                        row,
                        QUEST_AVAILABILITY_COLUMNS,
                        provenance=_prov(path, "LocalizationAvailability", row_number),
                    )
                )
    finally:
        wb.close()


def _ingest_four_language(path: Path, by_uid: dict[str, dict[str, Any]]) -> None:
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if "FourLanguageTable" in wb.sheetnames:
            for row_number, row in _rows_as_dicts(wb["FourLanguageTable"]):
                uid = _text(row.get("ContentUid"))
                if not uid:
                    continue
                record = _record(by_uid, uid)
                fields = REVIEW_COLUMNS + [
                    "TraditionalChineseReferenceType",
                    "TraditionalChineseReuseReference",
                    "SimplifiedChinesePresent",
                    "SimplifiedChineseVersion",
                    "SimplifiedChineseText",
                    "RussianPresent",
                    "RussianVersion",
                    "RussianText",
                ]
                record.setdefault("multilingualReview", []).append(
                    _project(row, fields, provenance=_prov(path, "FourLanguageTable", row_number))
                )
    finally:
        wb.close()


def _rows_as_dicts(sheet, *, required_headers: set[str] | None = None, max_header_rows: int = 10):
    required = required_headers or {"ContentUid"}
    rows = sheet.iter_rows(values_only=True)
    headers = None
    header_row_number = None
    for row_number in range(1, max_header_rows + 1):
        try:
            values = next(rows)
        except StopIteration:
            break
        candidate = [str(v).strip() if v is not None else "" for v in values]
        if required.issubset(set(candidate)):
            headers = candidate
            header_row_number = row_number
            break
    if headers is None or header_row_number is None:
        expected = ", ".join(sorted(required))
        raise ValueError(
            f"Sheet {sheet.title!r}: no trusted header found in rows 1-{max_header_rows}; "
            f"expected header signal(s): {expected}"
        )
    for row_number, values in enumerate(rows, start=header_row_number + 1):
        row = {header: value for header, value in zip(headers, values, strict=False) if header}
        if any(value not in (None, "") for value in row.values()):
            yield row_number, row


def _record(by_uid: dict[str, dict[str, Any]], uid: str) -> dict[str, Any]:
    return by_uid.setdefault(uid, {"contentUid": uid})


def _project(row: dict[str, Any], fields: list[str], *, provenance: dict[str, Any]) -> dict[str, Any]:
    result = {field: _json_cell(row.get(field)) for field in fields if row.get(field) not in (None, "")}
    result["provenance"] = provenance
    return result


def _prov(path: Path, sheet: str, row: int) -> dict[str, Any]:
    return {"workbook": path.name, "sheet": sheet, "row": row}


def _json_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _source_entry(path: Path) -> dict[str, Any]:
    return {"path": str(path), "name": path.name, "sha256": _sha256(path)}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        while chunk := fh.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _summarize(by_uid: dict[str, dict[str, Any]]) -> dict[str, Any]:
    records = [value for key, value in by_uid.items() if not key.startswith("__")]
    return {
        "contentUidCount": len(records),
        "withContextEvidence": sum(bool(r.get("contextEvidence")) for r in records),
        "withBarkEvidence": sum(bool(r.get("barkEvidence")) for r in records),
        "withBarkReview": sum(bool(r.get("barkReview")) for r in records),
        "withQuestReview": sum(bool(r.get("questReview")) for r in records),
        "withQuestLocalizationAvailability": sum(bool(r.get("questLocalizationAvailability")) for r in records),
        "withReviewEvidence": sum(bool(r.get("reviewEvidence")) for r in records),
        "withMultilingualReview": sum(bool(r.get("multilingualReview")) for r in records),
        "sharedSetMetadataPresent": "__shared_sets__" in by_uid,
    }


if __name__ == "__main__":
    raise SystemExit(main())
